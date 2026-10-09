"""
Migrate a Zalpora-format Excel into the 6-tab Google Sheet layout.

Usage:
    python migrate.py <path/to/zalpora.xlsx> <google_sheet_id>
    python migrate.py <path/to/zalpora.xlsx> --dry-run    (prints what it would write)

The target Google Sheet must already exist and be shared with the
service account as Editor. Any existing tabs with the six names below
WILL be overwritten.

The parsing logic is a line-for-line Python port of v32's
parseWorkbook() + extractTopNumber(), so a migrated village produces
the same shape the HTML used to see from an upload.
"""
from __future__ import annotations

import argparse
import re
import sys
from typing import Any, Dict, List, Optional

from io import BytesIO

from openpyxl import load_workbook


# -----------------------------------------------------------------------------
# Column indexes — mirrors the COL constant in v32
# -----------------------------------------------------------------------------
COL = {
    "STATUS":      0,
    "REMARKS":     1,
    "MUT_TYPE":    2,
    "MUT_NUMBER":  3,
    "DEMAND":      4,
    "RENT":        5,
    "QISM":        6,
    "MARLA":       7,
    "KANAL":       8,
    "KHASRA_NEW":  9,
    "KHASRA_OLD": 10,
    "IRRIGATION": 11,
    "CULTIVATOR": 12,
    "OWNER":      13,
    "PATTI":      14,
    "KHATA":      15,
    "KHEWAT":     16,
}


# -----------------------------------------------------------------------------
# Helpers (ported from v32)
# -----------------------------------------------------------------------------
def extract_top_number(cell: Any, mark_uncertain: bool = False) -> str:
    """Mirror of extractTopNumber() in v32."""
    if cell is None:
        return ""
    s = str(cell).strip()
    if s == "":
        return ""
    lines = [l.strip() for l in re.split(r"[\r\n]+", s)]
    lines = [l for l in lines if l and l != "-"]
    if not lines:
        return ""
    val = re.sub(r"\.+$", "", lines[0]).strip()
    if val == "":
        return ""
    if mark_uncertain and re.search(r"[/\\#*]", val):
        return "?" + val
    return val


def _s(v: Any) -> str:
    return "" if v is None else str(v).strip()


def _opt(v: Any) -> Optional[Any]:
    """Preserve numeric types where openpyxl gave them; None for empty."""
    if v is None:
        return None
    if isinstance(v, str) and v.strip() == "":
        return None
    return v


# -----------------------------------------------------------------------------
# Parser
# -----------------------------------------------------------------------------
def parse_workbook(xlsx_source) -> Dict[str, Any]:
    """Parse a Zalpora-format Excel. xlsx_source can be a path (str) OR
    bytes OR a file-like object (used by the admin page's upload handler)."""
    if isinstance(xlsx_source, (bytes, bytearray)):
        xlsx_source = BytesIO(xlsx_source)
    wb = load_workbook(xlsx_source, data_only=True, read_only=True)
    ws = wb["Jamabandi"] if "Jamabandi" in wb.sheetnames else wb[wb.sheetnames[0]]

    data: Dict[str, Any] = {"khewats": {}}
    current_khewat = None
    current_khata = None
    current_khasra = None
    khewat_order = 0
    khata_order_by_kew: Dict[str, int] = {}
    khasra_order_by_kt: Dict[tuple, int] = {}

    for row in ws.iter_rows(values_only=True):
        if row is None:
            continue
        # Pad short rows so index access is safe
        row = list(row) + [None] * (max(COL.values()) + 1 - len(row))

        qism = row[COL["QISM"]]
        status = row[COL["STATUS"]]
        khewat_raw = row[COL["KHEWAT"]]
        khata_raw = row[COL["KHATA"]]
        khasra_raw = row[COL["KHASRA_NEW"]]

        if not qism and not khewat_raw and not status:
            continue

        qism_str = str(qism).strip() if qism not in (None, "") else ""
        qlow = qism_str.lower()

        # Rollup / total rows
        if qlow.startswith("total area for khasra"):
            if current_khasra:
                current_khasra["totalKanal"] = row[COL["KANAL"]]
                current_khasra["totalMarla"] = row[COL["MARLA"]]
            continue
        if qlow.startswith("total area for khata"):
            if current_khata:
                current_khata["totalKanal"] = row[COL["KANAL"]]
                current_khata["totalMarla"] = row[COL["MARLA"]]
            continue
        if qlow.startswith("total area for khewat"):
            if current_khewat:
                current_khewat["totalKanal"] = row[COL["KANAL"]]
                current_khewat["totalMarla"] = row[COL["MARLA"]]
            continue

        # Sub-header
        if qism_str == "Qism Zameen":
            continue

        # Primary row
        if khewat_raw and qism_str:
            kew_no = extract_top_number(khewat_raw)
            kt_no = extract_top_number(khata_raw)
            kh_no = extract_top_number(khasra_raw, mark_uncertain=True)
            if not kew_no or not kt_no or not kh_no:
                continue

            if kew_no not in data["khewats"]:
                data["khewats"][kew_no] = {
                    "_no": kew_no,
                    "owner": "", "patti": "",
                    "totalKanal": None, "totalMarla": None,
                    "orderIndex": khewat_order,
                    "khatas": {},
                }
                khewat_order += 1
            current_khewat = data["khewats"][kew_no]

            if row[COL["OWNER"]] and not current_khewat["owner"]:
                current_khewat["owner"] = _s(row[COL["OWNER"]])
            if row[COL["PATTI"]] and not current_khewat["patti"]:
                current_khewat["patti"] = _s(row[COL["PATTI"]])

            if kt_no not in current_khewat["khatas"]:
                idx = khata_order_by_kew.get(kew_no, 0)
                khata_order_by_kew[kew_no] = idx + 1
                current_khewat["khatas"][kt_no] = {
                    "_no": kt_no,
                    "cultivator": "",
                    "totalKanal": None, "totalMarla": None,
                    "orderIndex": idx,
                    "khasras": {},
                }
            current_khata = current_khewat["khatas"][kt_no]

            if row[COL["CULTIVATOR"]] and not current_khata["cultivator"]:
                current_khata["cultivator"] = _s(row[COL["CULTIVATOR"]])

            key = (kew_no, kt_no)
            if kh_no not in current_khata["khasras"]:
                idx = khasra_order_by_kt.get(key, 0)
                khasra_order_by_kt[key] = idx + 1
                current_khata["khasras"][kh_no] = {
                    "_no": kh_no,
                    "portions": [],
                    "totalKanal": None, "totalMarla": None,
                    "mutationType":   _s(row[COL["MUT_TYPE"]]),
                    "mutationNumber": _s(row[COL["MUT_NUMBER"]]),
                    "demand":         _s(row[COL["DEMAND"]]),
                    "remarks":        _s(row[COL["REMARKS"]]),
                    "irrigation":     _s(row[COL["IRRIGATION"]]),
                    "khasraOld":      _s(row[COL["KHASRA_OLD"]]),
                    "status":         _s(status),
                    "orderIndex": idx,
                }
            current_khasra = current_khata["khasras"][kh_no]
            current_khasra["portions"].append({
                "qism": qism_str,
                "kanal": row[COL["KANAL"]] or 0,
                "marla": row[COL["MARLA"]] or 0,
            })
            continue

        # Continuation portion row (no khewat, has qism, current khasra open)
        if qism_str and current_khasra and not khewat_raw:
            current_khasra["portions"].append({
                "qism": qism_str,
                "kanal": row[COL["KANAL"]] or 0,
                "marla": row[COL["MARLA"]] or 0,
            })
            continue

    return data


# -----------------------------------------------------------------------------
# Flatten → six tabs
# -----------------------------------------------------------------------------
def flatten(parsed: Dict[str, Any]) -> Dict[str, List[list]]:
    """Return {tab_name: [[header_row], [row], [row], ...]}."""
    khewats_rows: List[list] = [
        ["khewat_no", "owner_raw", "patti", "total_kanal", "total_marla"]
    ]
    khatas_rows: List[list] = [
        ["khewat_no", "khata_no", "cultivator_raw", "total_kanal", "total_marla"]
    ]
    khasras_rows: List[list] = [
        ["khewat_no", "khata_no", "khasra_no_new", "khasra_no_old",
         "irrigation", "mutation_type", "mutation_number", "demand",
         "remarks", "status", "total_kanal", "total_marla"]
    ]
    portions_rows: List[list] = [
        ["khewat_no", "khata_no", "khasra_no", "portion_index",
         "qism", "kanal", "marla", "sqft"]
    ]
    mutations_rows: List[list] = [
        ["khewat_no", "khata_no", "khasra_no", "mutation_index",
         "mut_number", "mut_type"]
    ]

    # Sort by orderIndex throughout so sheet mirrors Excel order
    khewats = sorted(
        parsed["khewats"].items(), key=lambda kv: kv[1].get("orderIndex", 0)
    )
    for kew_no, kew in khewats:
        khewats_rows.append([
            kew_no, kew.get("owner", ""), kew.get("patti", ""),
            kew.get("totalKanal"), kew.get("totalMarla"),
        ])
        khatas = sorted(
            kew["khatas"].items(), key=lambda kv: kv[1].get("orderIndex", 0)
        )
        for kt_no, kt in khatas:
            khatas_rows.append([
                kew_no, kt_no, kt.get("cultivator", ""),
                kt.get("totalKanal"), kt.get("totalMarla"),
            ])
            khasras = sorted(
                kt["khasras"].items(), key=lambda kv: kv[1].get("orderIndex", 0)
            )
            for kh_no, kh in khasras:
                khasras_rows.append([
                    kew_no, kt_no, kh_no, kh.get("khasraOld", ""),
                    kh.get("irrigation", ""),
                    kh.get("mutationType", ""), kh.get("mutationNumber", ""),
                    kh.get("demand", ""), kh.get("remarks", ""),
                    kh.get("status", ""),
                    kh.get("totalKanal"), kh.get("totalMarla"),
                ])
                for pi, p in enumerate(kh.get("portions", [])):
                    portions_rows.append([
                        kew_no, kt_no, kh_no, pi,
                        p.get("qism", ""),
                        p.get("kanal") or 0,
                        p.get("marla") or 0,
                        "",  # sqft — new field, left blank on migration
                    ])
                # One mutation per khasra at migration time (if any)
                mt = kh.get("mutationType", "")
                mn = kh.get("mutationNumber", "")
                if mt or mn:
                    mutations_rows.append([
                        kew_no, kt_no, kh_no, 0, mn, mt,
                    ])

    meta_rows: List[list] = [
        ["key", "value"],
        ["village_name", ""],
        ["tehsil", ""],
        ["district", ""],
        ["halqa", ""],
        ["year", ""],
    ]

    return {
        "Khewats":   khewats_rows,
        "Khatas":    khatas_rows,
        "Khasras":   khasras_rows,
        "Portions":  portions_rows,
        "Mutations": mutations_rows,
        "_Meta":     meta_rows,
    }


# -----------------------------------------------------------------------------
# Writer
# -----------------------------------------------------------------------------
def write_sheets(sheet_id: str, tabs: Dict[str, List[list]]) -> None:
    """Thin shim — delegates to sheets.write_village(). Kept so CLI still works."""
    import sheets as sheets_mod  # local import so --dry-run doesn't need creds
    sheets_mod.write_village(sheet_id, tabs)


def migrate_bytes(xlsx_bytes: bytes, sheet_id: str) -> Dict[str, int]:
    """Full migrate from in-memory bytes, used by the admin page.
    Returns {'khewats': n, 'khatas': n, 'khasras': n, 'portions': n, 'mutations': n}.
    """
    parsed = parse_workbook(xlsx_bytes)
    tabs = flatten(parsed)
    import sheets as sheets_mod
    sheets_mod.write_village(sheet_id, tabs)
    return {
        "khewats":   len(tabs["Khewats"]) - 1,
        "khatas":    len(tabs["Khatas"]) - 1,
        "khasras":   len(tabs["Khasras"]) - 1,
        "portions":  len(tabs["Portions"]) - 1,
        "mutations": len(tabs["Mutations"]) - 1,
    }


# -----------------------------------------------------------------------------
# CLI
# -----------------------------------------------------------------------------
def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("xlsx", help="Path to the Zalpora-format Excel")
    ap.add_argument("sheet_id", nargs="?", default=None,
                    help="Target Google Sheet ID (omit for --dry-run)")
    ap.add_argument("--dry-run", action="store_true",
                    help="Print the tab counts without writing anything")
    args = ap.parse_args()

    parsed = parse_workbook(args.xlsx)
    tabs = flatten(parsed)

    n_khewats   = len(tabs["Khewats"]) - 1
    n_khatas    = len(tabs["Khatas"]) - 1
    n_khasras   = len(tabs["Khasras"]) - 1
    n_portions  = len(tabs["Portions"]) - 1
    n_mutations = len(tabs["Mutations"]) - 1

    print(f"Parsed {args.xlsx}:")
    print(f"  {n_khewats} khewats")
    print(f"  {n_khatas} khatas")
    print(f"  {n_khasras} khasras")
    print(f"  {n_portions} portions")
    print(f"  {n_mutations} mutations")

    if args.dry_run or not args.sheet_id:
        print("\n(dry run — nothing written)")
        return 0

    print(f"\nWriting to Google Sheet {args.sheet_id}...")
    write_sheets(args.sheet_id, tabs)
    print("Done. Verify the sheet, then share it with the user who should see it.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
